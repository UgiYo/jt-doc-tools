import io

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.tools.ppt_image_text_editor.editable_bridge import build_editable_deck

def test_image_ocr_becomes_native_text_overlay():
    image = Image.new("RGB", (400, 200), "white")
    ImageDraw.Draw(image).text((40, 50), "GitLab", fill="black")
    png = io.BytesIO(); image.save(png, "PNG")
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(1), Inches(8), Inches(4))
    source = io.BytesIO(); prs.save(source)
    analyses = [{"index": 0, "media_path": "ppt/media/image1.png", "width": 400, "height": 200,
                 "words": [{"left": 40, "top": 50, "width": 100, "height": 30, "text": "GitLab"}]}]
    deck = build_editable_deck(source.getvalue(), analyses, [])
    elements = deck["slides"][0]["elements"]
    assert [e for e in elements if e["type"] == "image"]
    assert any(e["type"] == "text" and e["text"] == "GitLab" and e["source"] == "ocr" for e in elements)
