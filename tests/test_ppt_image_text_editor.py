import importlib.util
import io
import pathlib
import zipfile
import sys
from PIL import Image, ImageDraw

ROOT = pathlib.Path(__file__).resolve().parents[1]

def load(name, rel):
    spec = importlib.util.spec_from_file_location(name, ROOT / rel)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod

image_edit = load('pite_image_edit', 'app/tools/ppt_image_text_editor/image_edit.py')
pptx_core = load('pite_pptx_core', 'app/tools/ppt_image_text_editor/pptx_core.py')
edit_text = image_edit.edit_text
list_slide_images = pptx_core.list_slide_images
read_media = pptx_core.read_media
replace_media = pptx_core.replace_media


def _fake_pptx(img: bytes) -> bytes:
    slide = b'''<?xml version="1.0"?><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:cSld><p:spTree><p:pic><p:blipFill><a:blip r:embed="rId2"/></p:blipFill></p:pic></p:spTree></p:cSld></p:sld>'''
    rels = b'''<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>'''
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as z:
        z.writestr('ppt/slides/slide1.xml',slide)
        z.writestr('ppt/slides/_rels/slide1.xml.rels',rels)
        z.writestr('ppt/media/image1.png',img)
    return out.getvalue()


def test_list_and_replace_media_preserves_relationships():
    im=Image.new('RGB',(100,40),'white'); b=io.BytesIO(); im.save(b,'PNG')
    pptx=_fake_pptx(b.getvalue())
    refs=list_slide_images(pptx)
    assert len(refs)==1 and refs[0].slide==1 and refs[0].media_path=='ppt/media/image1.png'
    changed=Image.new('RGB',(100,40),'red'); c=io.BytesIO(); changed.save(c,'PNG')
    out=replace_media(pptx,{'ppt/media/image1.png':c.getvalue()})
    assert read_media(out,'ppt/media/image1.png')==c.getvalue()
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert b'rId2' in z.read('ppt/slides/slide1.xml')


def test_edit_text_keeps_canvas_size():
    im=Image.new('RGB',(200,80),'white'); d=ImageDraw.Draw(im); d.text((20,20),'Gitlab',fill='black'); b=io.BytesIO(); im.save(b,'PNG')
    edited=edit_text(b.getvalue(),box=(18,18,95,45),new_text='GitLab')
    out=Image.open(io.BytesIO(edited))
    assert out.size==(200,80)


def test_image_editors_are_registered_for_the_ui():
    from app.tool_registry import discover_tools

    ids = {tool.metadata.id for tool in discover_tools()}
    assert {"ppt-image-text-editor", "image-text-editor"} <= ids


def test_standalone_image_editor_applies_edits():
    from app.tools.image_text_editor.router import _apply

    image = Image.new("RGB", (200, 80), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 20), "Before", fill="black")
    source = io.BytesIO()
    image.save(source, "PNG")

    edited = _apply(source.getvalue(), [{
        "left": 18,
        "top": 18,
        "width": 80,
        "height": 30,
        "new_text": "After",
    }])

    result = Image.open(io.BytesIO(edited))
    assert result.size == (200, 80)
    assert result.format == "PNG"


def test_editable_pptx_download_contains_native_text_box():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
    from app.tools.editable_slides.pptx_io import export_pptx
    from app.tools.ppt_image_text_editor.editable_bridge import build_editable_deck

    image = Image.new("RGB", (1280, 720), "white")
    ImageDraw.Draw(image).text((120, 100), "Original OCR text", fill="black")
    png = io.BytesIO()
    image.save(png, "PNG")
    source_prs = Presentation()
    source_prs.slide_width = Inches(13.333)
    source_prs.slide_height = Inches(7.5)
    source_slide = source_prs.slides.add_slide(source_prs.slide_layouts[6])
    source_slide.shapes.add_picture(io.BytesIO(png.getvalue()), 0, 0, source_prs.slide_width, source_prs.slide_height)
    source = io.BytesIO()
    source_prs.save(source)
    ref = list_slide_images(source.getvalue())[0]

    analyses = [{
        "media_path": ref.media_path, "slides": [1], "rel_ids": [ref.rel_id],
        "index": 0, "width": 1280, "height": 720,
        "words": [{"left": 120, "top": 100, "width": 260, "height": 45,
                   "text": "Original OCR text", "conf": 99}],
    }]
    edits = [{
        "image_index": 0, "word_index": 0, "left": 120, "top": 100,
        "width": 260, "height": 45, "old_text": "Original OCR text",
        "new_text": "可編輯文字", "font_family": "Microsoft JhengHei", "font_size": 24,
    }]
    download = export_pptx(build_editable_deck(source.getvalue(), analyses, edits))

    assert download[:2] == b"PK"
    result = Presentation(io.BytesIO(download))
    assert len(result.slides) == 1
    texts = [shape.text for shape in result.slides[0].shapes if getattr(shape, "has_text_frame", False)]
    pictures = [shape for shape in result.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert "可編輯文字" in texts
    assert len(pictures) == 1


def test_ocr_fragments_are_grouped_into_fitted_line_text():
    from app.tools.ppt_image_text_editor.editable_bridge import _line_segments, _segment_text
    words = [
        {"left": 10, "top": 10, "width": 30, "height": 20, "text": "Dev"},
        {"left": 43, "top": 11, "width": 25, "height": 19, "text": "Ops"},
        {"left": 300, "top": 10, "width": 30, "height": 20, "text": "Other"},
        {"left": 10, "top": 50, "width": 30, "height": 20, "text": "Next"},
    ]
    segments = _line_segments(words)
    assert len(segments) == 3
    assert _segment_text(segments[0], {}, 0) == "DevOps"


def test_text_removal_mask_covers_entire_expanded_ocr_box():
    from PIL import Image
    from app.tools.ppt_image_text_editor.editable_bridge import _text_mask

    image = Image.new("RGB", (200, 100), "white")
    mask = _text_mask(image, [{"left": 50, "top": 30, "width": 80, "height": 20, "text": "Text"}])
    assert mask[30:50, 50:130].min() == 255
    assert mask[27:53, 47:133].max() == 255
    assert mask[0, 0] == 0


def test_cleaned_background_has_no_dark_text_fragments_and_preserves_neighbours():
    import numpy as np
    from app.tools.ppt_image_text_editor.editable_bridge import _clean_image

    image = Image.new("RGB", (400, 160), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 60), "Old text", fill="black")
    draw.ellipse((20, 50, 70, 100), fill=(230, 55, 90))
    source = io.BytesIO()
    image.save(source, "PNG")
    cleaned = Image.open(io.BytesIO(_clean_image(source.getvalue(), [
        {"left": 98, "top": 57, "width": 90, "height": 22, "text": "Old text"}
    ]))).convert("RGB")
    pixels = np.asarray(cleaned)
    assert ((pixels[57:82, 98:190].mean(axis=2)) < 100).sum() == 0
    assert tuple(pixels[75, 45]) == (230, 55, 90)


def test_export_sets_cjk_typeface_and_keeps_fixed_font_size():
    import zipfile
    from app.tools.editable_slides.pptx_io import export_pptx

    deck = {"size": {"width": 13.333, "height": 7.5}, "slides": [{
        "background": "#ffffff", "elements": [{
            "type": "text", "x": 1, "y": 1, "width": 5, "height": .5,
            "text": "繁體中文 DevOps", "fontFamily": "Microsoft JhengHei",
            "fontSize": 24, "source": "ocr", "fitText": False,
        }]
    }]}
    raw = export_pptx(deck)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert '<a:ea typeface="Microsoft JhengHei"' in xml
    assert '<a:latin typeface="Microsoft JhengHei"' in xml
    assert 'lang="zh-TW"' in xml
    assert "normAutofit" not in xml and "spAutoFit" not in xml
    assert 'sz="2400"' in xml
