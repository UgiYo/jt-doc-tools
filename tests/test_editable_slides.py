import io

from pptx import Presentation

from app.tools.editable_slides.pptx_io import export_pptx, import_pptx
from app.tools.editable_slides.schema import DEFAULT_DECK, normalize_deck

def test_editable_slide_round_trip():
    raw = export_pptx(normalize_deck(DEFAULT_DECK))
    prs = Presentation(io.BytesIO(raw))
    assert len(prs.slides) == 1
    assert any("可編輯簡報" in shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
    imported = import_pptx(raw)
    assert imported["slides"][0]["elements"]

def test_editable_slides_registered():
    from app.tool_registry import discover_tools
    assert "editable-slides" in {tool.metadata.id for tool in discover_tools()}
