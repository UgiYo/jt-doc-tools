from __future__ import annotations

import base64
import io
import uuid

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

EMU_PER_INCH = 914400

def _rgb(value: str | None, fallback="000000") -> RGBColor:
    value = (value or fallback).lstrip("#")
    return RGBColor.from_string((value if len(value) == 6 else fallback).upper())

def _inch(value): return round(value / EMU_PER_INCH, 4)

def _set_text(shape, el):
    tf = shape.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if el.get("source") == "ocr":
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        if el.get("fitText"): tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    else:
        tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(el.get("align"), PP_ALIGN.CENTER)
    run = p.add_run(); run.text = str(el.get("text", ""))
    run.font.name = el.get("fontFamily", "Microsoft JhengHei")
    run.font.size = Pt(float(el.get("fontSize", 20))); run.font.bold = bool(el.get("bold", False))
    run.font.italic = bool(el.get("italic", False)); run.font.underline = bool(el.get("underline", False))
    run.font.color.rgb = _rgb(el.get("color"), "111827")

def export_pptx(deck: dict) -> bytes:
    prs = Presentation(); prs.slide_width = Inches(deck["size"]["width"]); prs.slide_height = Inches(deck["size"]["height"])
    for source in deck["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6]); bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = _rgb(source.get("background"), "FFFFFF")
        for el in source["elements"]:
            pos = [Inches(el[k]) for k in ("x", "y", "width", "height")]
            if el["type"] == "text":
                _set_text(slide.shapes.add_textbox(*pos), el)
            elif el["type"] == "shape":
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *pos)
                shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(el.get("fill"), "E5E7EB")
                shape.line.color.rgb = _rgb(el.get("border") or el.get("fill"), "E5E7EB"); _set_text(shape, el)
            else:
                src = el.get("src", "")
                if src.startswith("data:image/") and ";base64," in src:
                    payload = base64.b64decode(src.split(",", 1)[1], validate=True)
                    slide.shapes.add_picture(io.BytesIO(payload), *pos)
    out = io.BytesIO(); prs.save(out); return out.getvalue()

def import_pptx(raw: bytes) -> dict:
    prs = Presentation(io.BytesIO(raw)); slides = []
    for index, slide in enumerate(prs.slides):
        elements = []
        for shape in slide.shapes:
            base = {"id": uuid.uuid4().hex, "x": _inch(shape.left), "y": _inch(shape.top),
                    "width": _inch(shape.width), "height": _inch(shape.height)}
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                base.update(type="image", src=f"data:{image.content_type};base64," + base64.b64encode(image.blob).decode(), ocrCandidate=True)
                elements.append(base)
            elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                run = next((r for p in shape.text_frame.paragraphs for r in p.runs), None)
                base.update(type="shape" if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else "text", text=shape.text,
                            fontSize=round(run.font.size.pt if run and run.font.size else 20, 1),
                            fontFamily=(run.font.name if run and run.font.name else "Microsoft JhengHei"),
                            bold=bool(run.font.bold) if run else False, italic=bool(run.font.italic) if run else False,
                            underline=bool(run.font.underline) if run else False, color="#111827", fill="#E5E7EB")
                elements.append(base)
        slides.append({"id": f"slide-{index + 1}", "background": "#FFFFFF", "elements": elements})
    return {"version": 1, "title": "Imported presentation", "size": {"width": _inch(prs.slide_width), "height": _inch(prs.slide_height)}, "slides": slides}
